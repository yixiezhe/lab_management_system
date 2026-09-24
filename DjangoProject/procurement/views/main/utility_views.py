import io
import zipfile
import os
import tempfile
from datetime import datetime
from decimal import Decimal, InvalidOperation
from urllib.parse import quote
from django.http import HttpResponse, Http404

import xlrd
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from rest_framework import status
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView
from django.db import transaction
from django.conf import settings
from django.utils import timezone

from ...models import Platform, PurchaseRequest, RequestItem
from ...permissions import IsApprover, IsSystemAdmin  # <-- 导入 IsSystemAdmin
from common_serializers.serializers import PurchaseRequestSerializer


# ==========================================================================
# 辅助函数
# ==========================================================================
def _format_total_price(value):
    if value is None:
        return "0.00"
    try:
        return f"{Decimal(value):.2f}"
    except (InvalidOperation, TypeError):
        return str(value)


def _build_attachment_name(total_price, label, index=None):
    total_text = _format_total_price(total_price)
    suffix = f"_{label}"
    if index is not None:
        suffix = f"{suffix}_{index}"
    return f"{total_text}{suffix}"


def _add_file_to_zip(zip_file, obj, field_name, new_name_prefix):
    """ 辅助函数：检查文件是否存在并添加到ZIP包 """
    if obj and hasattr(obj, field_name):
        file_field = getattr(obj, field_name)
        if file_field and hasattr(file_field, 'path') and os.path.exists(file_field.path):
            _, ext = os.path.splitext(file_field.name)
            new_name = f"{new_name_prefix}{ext}"
            zip_file.write(file_field.path, new_name)

def _add_invoices_to_zip(zip_file, invoices, total_price):
    invoice_list = list(invoices) if invoices is not None else []
    if not invoice_list:
        return
    for index, invoice in enumerate(invoice_list, start=1):
        name_prefix = _build_attachment_name(total_price, '发票', index if len(invoice_list) > 1 else None)
        _add_file_to_zip(zip_file, invoice, 'invoice_image', name_prefix)


def _add_reimbursement_to_zip(zip_file, reimbursement, total_price):
    if not reimbursement:
        return
    name_prefix = _build_attachment_name(total_price, '报销单')
    _add_file_to_zip(zip_file, reimbursement, 'reimbursement_photo', name_prefix)


def _find_invoice_template_path():
    base_dir = settings.BASE_DIR.parent
    base_bytes = os.fsencode(str(base_dir))
    xlsx_candidates = []
    xls_candidates = []
    for entry in os.listdir(base_bytes):
        decoded_entry = os.fsdecode(entry).lower()
        if '发票' not in decoded_entry and 'invoice' not in decoded_entry:
            continue
        lower_entry = entry.lower()
        if lower_entry.endswith(b'.xlsx'):
            xlsx_candidates.append(os.path.join(base_bytes, entry))
        elif lower_entry.endswith(b'.xls'):
            xls_candidates.append(os.path.join(base_bytes, entry))
    xlsx_candidates.sort(key=lambda path: os.fsdecode(os.path.basename(path)).lower())
    xls_candidates.sort(key=lambda path: os.fsdecode(os.path.basename(path)).lower())
    if xlsx_candidates:
        return xlsx_candidates[0]
    if xls_candidates:
        return xls_candidates[0]
    return None


def _fill_invoice_rows(sheet, rows, start_row=2):
    for row_index, (invoice_number, total_price) in enumerate(rows, start_row):
        cell_invoice = sheet.cell(row=row_index, column=1)
        cell_invoice.number_format = "@"
        cell_invoice.value = str(invoice_number) if invoice_number is not None else ''
        sheet.cell(row=row_index, column=2).value = float(total_price) if total_price is not None else ''


def _build_invoice_excel(rows, total_sum):
    template_path_bytes = _find_invoice_template_path()
    output_name = f"{_format_total_price(total_sum)}发票信息.xlsx"

    with tempfile.TemporaryDirectory() as tmp_dir:
        output_path = os.path.join(tmp_dir, output_name)
        if template_path_bytes:
            template_path = os.fsdecode(template_path_bytes)
            template_ext = os.path.splitext(template_path)[1].lower()

            if template_ext == '.xlsx':
                workbook = load_workbook(template_path)
                sheet = workbook.active
                _fill_invoice_rows(sheet, rows, start_row=2)
                workbook.save(output_path)
            else:
                template_book = xlrd.open_workbook(template_path)
                template_sheet = template_book.sheet_by_index(0)

                workbook = Workbook()
                sheet = workbook.active
                sheet.title = template_sheet.name

                # Copy header row and column widths from the template.
                if template_sheet.nrows > 0:
                    for col_idx in range(template_sheet.ncols):
                        header_value = template_sheet.cell_value(0, col_idx)
                        sheet.cell(row=1, column=col_idx + 1, value=header_value)
                        col_info = template_sheet.colinfo_map.get(col_idx)
                        if col_info:
                            sheet.column_dimensions[get_column_letter(col_idx + 1)].width = col_info.width / 256.0

                _fill_invoice_rows(sheet, rows, start_row=2)
                workbook.save(output_path)
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = '发票信息'
            sheet.cell(row=1, column=1, value='发票号码')
            sheet.cell(row=1, column=2, value='金额')
            sheet.column_dimensions['A'].width = 28
            sheet.column_dimensions['B'].width = 18
            _fill_invoice_rows(sheet, rows, start_row=2)
            workbook.save(output_path)

        with open(output_path, 'rb') as file_handle:
            return output_name, file_handle.read()


C2C_RECORD_HEADERS_BASE = [
    '序号', '申请ID', '单号', '申请日期', '当前状态', '平台', '申请人', '导师', '采购人员', '审批人',
    '采购明细', 'CAS号', '厂商', '货号', '采购链接', '申请总金额', '实际支付金额',
    '收货照片', '发票单号', '发票公司', '付款链接', '收货情况', '支付时间', '验收单号', '报销单号', '报销实际金额',
    '请购单文件', '合同文件', '支付截图'
]
C2C_RECORD_HEADERS_AFTER_INVOICE = ['验收照片', '报销单照片']
C2C_IMAGE_EXTENSIONS = {'.jpg', '.jpeg', '.png', '.gif', '.bmp'}


def _join_unique(values, separator='、'):
    seen = []
    for value in values:
        if value is None:
            continue
        text = str(value).strip()
        if text and text not in seen:
            seen.append(text)
    return separator.join(seen)


def _money_value(value):
    if value is None or value == '':
        return None
    try:
        return float(Decimal(value))
    except (InvalidOperation, TypeError, ValueError):
        return value


def _format_date(value):
    if not value:
        return ''
    return value.strftime('%Y-%m-%d') if hasattr(value, 'strftime') else str(value)


def _format_datetime(value):
    if not value:
        return ''
    if hasattr(value, 'strftime'):
        return timezone.localtime(value).strftime('%Y-%m-%d %H:%M:%S') if timezone.is_aware(value) else value.strftime('%Y-%m-%d %H:%M:%S')
    return str(value)


def _request_children(request_obj):
    manager = getattr(request_obj, 'merged_children', None)
    if not manager:
        return []
    try:
        return list(manager.all())
    except Exception:
        return []


def _request_group(request_obj):
    return [request_obj] + _request_children(request_obj)


def _get_related(obj, name):
    try:
        return getattr(obj, name)
    except Exception:
        return None


def _collect_related(request_obj, name):
    related = []
    for item in _request_group(request_obj):
        value = _get_related(item, name)
        if value:
            related.append(value)
    return related


def _collect_invoices(request_obj):
    invoices = []
    for item in _request_group(request_obj):
        manager = getattr(item, 'invoices', None)
        if not manager:
            continue
        try:
            invoices.extend(list(manager.all()))
        except Exception:
            pass
    return sorted(invoices, key=lambda inv: (inv.created_at is None, inv.created_at or datetime.min, inv.id or 0))


def _is_rejected_request(request_obj):
    return getattr(request_obj, 'status', None) == 'rejected'


def _c2c_platform_label(platform_name=None):
    name = (platform_name or '选中').strip()
    return name if name.endswith('平台') else f'{name}平台'


def _c2c_export_title(platform_name, suffix):
    return f"公对公{_c2c_platform_label(platform_name)}{suffix}"


def _applicant_names(request_obj):
    children = _request_children(request_obj)
    if children:
        return _join_unique(child.applicant.name for child in children if getattr(child, 'applicant', None))
    return request_obj.applicant.name if getattr(request_obj, 'applicant', None) else ''


def _tutor_names(request_obj):
    children = _request_children(request_obj)
    source = children if children else [request_obj]
    return _join_unique(
        item.applicant.assigned_tutor.name
        for item in source
        if getattr(item, 'applicant', None) and getattr(item.applicant, 'assigned_tutor', None)
    )


def _handler_names(request_obj):
    handlers = [request_obj.handler] if getattr(request_obj, 'handler', None) else []
    handlers.extend(child.handler for child in _request_children(request_obj) if getattr(child, 'handler', None))
    return _join_unique(handler.name for handler in handlers if handler)


def _approval_names(request_obj):
    approvers = [request_obj.approved_by] if getattr(request_obj, 'approved_by', None) else []
    approvers.extend(child.approved_by for child in _request_children(request_obj) if getattr(child, 'approved_by', None))
    return _join_unique(approver.name for approver in approvers if approver)


def _purchase_detail(request_obj):
    details = []
    for item in request_obj.items.all():
        subtotal = Decimal(item.unit_price or 0) * Decimal(item.quantity or 0)
        spec = item.specifications or item.parameters or item.manufacturer or ''
        parts = [item.content]
        if spec:
            parts.append(spec)
        parts.append(f"{_format_total_price(item.unit_price)} x {item.quantity}")
        parts.append(f"小计 {_format_total_price(subtotal)}")
        details.append(' | '.join(parts))
    return '\n'.join(details)


def _purchase_item_values(request_obj, field_name):
    return _join_unique(getattr(item, field_name, None) for item in request_obj.items.all())


def _purchase_links(request_obj):
    return '\n'.join(
        link for link in _join_unique((item.purchase_link for item in request_obj.items.all()), '\n').split('\n') if link
    )


def _first_file_field(objects, field_name):
    for obj in objects:
        file_field = getattr(obj, field_name, None)
        if file_field:
            return file_field
    return None


def _attachment_label(file_field):
    if not file_field:
        return ''
    if hasattr(file_field, 'path') and os.path.exists(file_field.path):
        ext = os.path.splitext(file_field.name or '')[1].lower()
        return '见图' if ext in C2C_IMAGE_EXTENSIONS else '见文件'
    return '文件缺失'


def _add_file_preview(sheet, cell, file_field):
    if not file_field or not hasattr(file_field, 'path') or not os.path.exists(file_field.path):
        return

    ext = os.path.splitext(file_field.name or '')[1].lower()
    if ext not in C2C_IMAGE_EXTENSIONS:
        return

    try:
        image = ExcelImage(file_field.path)
        max_width, max_height = 120, 80
        scale = min(max_width / image.width, max_height / image.height, 1)
        image.width = int(image.width * scale)
        image.height = int(image.height * scale)
        sheet.add_image(image, cell.coordinate)
        sheet.row_dimensions[cell.row].height = max(sheet.row_dimensions[cell.row].height or 18, 66)
    except Exception:
        return


def _set_cell_value(sheet, row_number, col_number, value):
    cell = sheet.cell(row=row_number, column=col_number)
    cell.value = value
    cell.alignment = Alignment(vertical='center', wrap_text=True)
    return cell


def _build_c2c_record_rows(requests_to_export):
    rows = []
    max_invoice_count = 1
    for index, req in enumerate(requests_to_export, start=1):
        invoices = _collect_invoices(req)
        max_invoice_count = max(max_invoice_count, len(invoices))
        acceptances = _collect_related(req, 'acceptance')
        payments = _collect_related(req, 'payment')
        inspections = _collect_related(req, 'inspection')
        reimbursements = _collect_related(req, 'reimbursement_details')
        purchase_orders = _collect_related(req, 'purchase_order')
        contracts = _collect_related(req, 'contract')
        is_rejected = _is_rejected_request(req)

        rows.append({
            'index': index,
            'request': req,
            'invoices': invoices,
            'values': [
                index,
                req.id,
                req.order_number or '',
                _format_date(req.request_date),
                req.get_status_display(),
                req.platform,
                _applicant_names(req),
                _tutor_names(req),
                _handler_names(req),
                _approval_names(req),
                _purchase_detail(req),
                _purchase_item_values(req, 'cas_number'),
                _purchase_item_values(req, 'manufacturer'),
                _purchase_item_values(req, 'product_number'),
                _purchase_links(req),
                None if is_rejected else _money_value(req.total_price),
                None if is_rejected else _money_value(req.actual_payment_amount),
                _join_unique(inv.invoice_number for inv in invoices),
                _join_unique(inv.company_name for inv in invoices),
                '\n'.join(link for link in _join_unique((inv.payment_link for inv in invoices), '\n').split('\n') if link),
                _join_unique(acc.receiving_status for acc in acceptances),
                _join_unique(_format_datetime(pay.paid_at) for pay in payments),
                _join_unique(ins.inspection_number for ins in inspections),
                _join_unique(reim.reimbursement_number for reim in reimbursements),
                _money_value(sum((reim.actual_amount or Decimal('0.00')) for reim in reimbursements)) if reimbursements else None,
            ],
            'attachments': {
                'purchase_order': _first_file_field(purchase_orders, 'po_form'),
                'contract': _first_file_field(contracts, 'contract_form'),
                'payment': _first_file_field(payments, 'payment_photo'),
                'acceptance': _first_file_field(acceptances, 'acceptance_photo'),
                'inspection': _first_file_field(inspections, 'inspection_photo'),
                'reimbursement': _first_file_field(reimbursements, 'reimbursement_photo'),
            },
        })
    return rows, max_invoice_count


def _build_c2c_records_excel(requests_to_export, platform_name=None):
    rows, max_invoice_count = _build_c2c_record_rows(requests_to_export)
    invoice_headers = [f'发票图片{i}' for i in range(1, max_invoice_count + 1)]
    headers = C2C_RECORD_HEADERS_BASE + invoice_headers + C2C_RECORD_HEADERS_AFTER_INVOICE

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '购买记录'
    summary = workbook.create_sheet('金额汇总')

    title = _c2c_export_title(platform_name, '购买记录')
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = sheet.cell(row=1, column=1, value=title)
    title_cell.font = Font(size=16, bold=True)
    title_cell.alignment = Alignment(horizontal='center', vertical='center')

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    subtitle_cell = sheet.cell(
        row=2,
        column=1,
        value=f"来源：实验室管理系统；生成时间：{timezone.localtime().strftime('%Y-%m-%d %H:%M:%S')}"
    )
    subtitle_cell.alignment = Alignment(horizontal='left', vertical='center')

    header_fill = PatternFill('solid', fgColor='D9EAF7')
    thin = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=3, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    widths = [8, 10, 22, 13, 16, 12, 12, 12, 12, 12, 45, 18, 22, 18, 45, 14, 14, 22, 28, 24, 35, 35, 20, 20, 20, 14, 24, 24, 22]
    widths.extend([22] * max_invoice_count)
    widths.extend([22, 22])
    for col, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width

    total_sum = Decimal('0.00')
    actual_sum = Decimal('0.00')
    export_sum = Decimal('0.00')
    status_summary = {}
    embedded_images = 0
    missing_files = 0

    for row_offset, row_data in enumerate(rows, start=4):
        req = row_data['request']
        status_label = req.get_status_display()
        status_data = status_summary.setdefault(status_label, {
            'count': 0,
            'total_sum': Decimal('0.00'),
            'export_sum': Decimal('0.00'),
            'has_amount': False,
        })
        status_data['count'] += 1
        is_rejected = _is_rejected_request(req)

        if not is_rejected:
            try:
                request_total = Decimal(req.total_price or 0)
                total_sum += request_total
                status_data['total_sum'] += request_total
                status_data['has_amount'] = True
            except (InvalidOperation, TypeError):
                request_total = Decimal('0.00')
            if req.actual_payment_amount is not None:
                try:
                    actual_amount = Decimal(req.actual_payment_amount)
                    actual_sum += actual_amount
                    export_sum += actual_amount
                    status_data['export_sum'] += actual_amount
                    status_data['has_amount'] = True
                except (InvalidOperation, TypeError):
                    pass
            else:
                try:
                    export_sum += request_total
                    status_data['export_sum'] += request_total
                    status_data['has_amount'] = True
                except (InvalidOperation, TypeError):
                    pass

        attachments = row_data['attachments']
        values = list(row_data['values'][:17])
        values.append(_attachment_label(attachments['acceptance']))
        values.extend(row_data['values'][17:])
        values.extend([
            _attachment_label(attachments['purchase_order']),
            _attachment_label(attachments['contract']),
            _attachment_label(attachments['payment']),
        ])

        invoices = row_data['invoices']
        for invoice_index in range(max_invoice_count):
            invoice = invoices[invoice_index] if invoice_index < len(invoices) else None
            values.append(_attachment_label(invoice.invoice_image if invoice else None))

        values.extend([
            _attachment_label(attachments['inspection']),
            _attachment_label(attachments['reimbursement']),
        ])

        for col, value in enumerate(values, start=1):
            cell = _set_cell_value(sheet, row_offset, col, value if value is not None else '')
            cell.border = border
            if col in (16, 17, 26):
                cell.number_format = '#,##0.00'

        file_columns = {
            18: attachments['acceptance'],
            27: attachments['purchase_order'],
            28: attachments['contract'],
            29: attachments['payment'],
        }
        invoice_start_col = 30
        for invoice_index, invoice in enumerate(invoices, start=0):
            file_columns[invoice_start_col + invoice_index] = invoice.invoice_image
        file_columns[invoice_start_col + max_invoice_count] = attachments['inspection']
        file_columns[invoice_start_col + max_invoice_count + 1] = attachments['reimbursement']

        for col, file_field in file_columns.items():
            if file_field and hasattr(file_field, 'path') and not os.path.exists(file_field.path):
                missing_files += 1
            before_count = len(sheet._images)
            _add_file_preview(sheet, sheet.cell(row=row_offset, column=col), file_field)
            embedded_images += len(sheet._images) - before_count

    sheet.freeze_panes = 'A4'

    summary.title = '金额汇总'
    summary_title = _c2c_export_title(platform_name, '金额汇总')
    summary_rows = [
        [summary_title, None, None, None],
        ['项目', '值', None, None],
        ['记录数', len(rows), None, None],
        ['申请总金额合计', float(total_sum), None, None],
        ['已填写实际支付金额合计', float(actual_sum), None, None],
        ['导出口径金额合计（有实际支付金额时优先，否则用申请总金额）', float(export_sum), None, None],
        ['嵌入图片数量', embedded_images, None, None],
        ['数据库有记录但本地文件不存在数量', missing_files, None, None],
        [None, None, None, None],
        [None, None, None, None],
        ['按状态汇总', None, None, None],
        ['状态', '记录数', '申请总金额', '导出口径金额'],
    ]
    summary_rows.extend(
        [
            status_label,
            data['count'],
            float(data['total_sum']) if data['has_amount'] else None,
            float(data['export_sum']) if data['has_amount'] else None,
        ]
        for status_label, data in status_summary.items()
    )
    summary.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    for row_number, row_values in enumerate(summary_rows, start=1):
        for col_number, value in enumerate(row_values, start=1):
            cell = summary.cell(row=row_number, column=col_number, value=value)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if row_number == 1:
                cell.font = Font(size=14, bold=True)
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            if row_number in (2, 12):
                cell.font = Font(bold=True)
                cell.fill = header_fill
            if row_number in (11,):
                cell.font = Font(bold=True)
            if (col_number == 2 and row_number in (4, 5, 6)) or (col_number in (3, 4) and row_number >= 13):
                cell.number_format = '#,##0.00'
            if row_number not in (9, 10):
                cell.border = border
    summary.column_dimensions['A'].width = 48
    summary.column_dimensions['B'].width = 18
    summary.column_dimensions['C'].width = 18
    summary.column_dimensions['D'].width = 18

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    date_text = timezone.localtime().strftime('%Y%m%d')
    filename = f"{title}_{date_text}.xlsx"
    return filename, output.getvalue()


class DashboardStatsAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, *args, **kwargs):
        user = self.request.user
        stats = {}
        my_pending_requests = PurchaseRequest.objects.filter(applicant=user, status__in=['pending', 'approved', 'paid',
                                                                                         'goods_received', 'invoiced',
                                                                                         'accepted',
                                                                                         'inspection_skipped']).count()
        stats['my_pending_requests'] = my_pending_requests

        is_approver_perm = IsApprover()
        if is_approver_perm.has_permission(request, self):
            awaiting_approval_count = PurchaseRequest.objects.filter(status__in=['pending', 'payment_rejected']).count()
            stats['awaiting_approval_count'] = awaiting_approval_count

        return Response(stats, status=status.HTTP_200_OK)


class PackageReceiptsView(APIView):
    permission_classes = [IsAuthenticated, IsApprover]

    def post(self, request, *args, **kwargs):
        request_ids = request.data.get('request_ids', [])
        if not request_ids or not isinstance(request_ids, list):
            return Response({"error": "请提供需要打包的申请ID列表。"}, status=status.HTTP_400_BAD_REQUEST)

        requests_to_package_qs = PurchaseRequest.objects.filter(id__in=request_ids)
        if requests_to_package_qs.count() != len(set(request_ids)):
            return Response({"error": "一个或多个申请ID无效。"}, status=status.HTTP_400_BAD_REQUEST)

        for req in requests_to_package_qs:
            self.check_object_permissions(request, req)

        requests_to_package = requests_to_package_qs.prefetch_related(
            'merged_children__payment', 'merged_children__acceptance', 'merged_children__invoices',
            'merged_children__inspection', 'merged_children__purchase_order', 'merged_children__contract',
            'invoices'
        ).select_related(
            'payment', 'acceptance', 'inspection', 'purchase_order', 'contract'
        )

        buffer = io.BytesIO()
        invoice_rows = []
        selected_total_sum = Decimal('0.00')
        with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            for req in requests_to_package:
                try:
                    selected_total_sum += Decimal(req.total_price or 0)
                except (InvalidOperation, TypeError):
                    pass
                if hasattr(req, 'merged_children') and req.merged_children.exists():
                    if hasattr(req, 'purchase_order'): _add_file_to_zip(
                        zip_file, req.purchase_order, 'po_form',
                        _build_attachment_name(req.total_price, '请购单')
                    )
                    if hasattr(req, 'contract'): _add_file_to_zip(
                        zip_file, req.contract, 'contract_form',
                        _build_attachment_name(req.total_price, '合同')
                    )
                    if hasattr(req, 'inspection'): _add_file_to_zip(
                        zip_file, req.inspection, 'inspection_photo',
                        _build_attachment_name(req.total_price, '验收凭证')
                    )

                    for child_req in req.merged_children.all():
                        if hasattr(child_req, 'payment'): _add_file_to_zip(
                            zip_file, child_req.payment, 'payment_photo',
                            _build_attachment_name(child_req.total_price, '支付凭证')
                        )
                        if hasattr(child_req, 'acceptance'): _add_file_to_zip(
                            zip_file, child_req.acceptance, 'acceptance_photo',
                            _build_attachment_name(child_req.total_price, '收货凭证')
                        )
                        _add_invoices_to_zip(zip_file, child_req.invoices.all(), child_req.total_price)
                        for invoice in child_req.invoices.all():
                            invoice_rows.append((invoice.invoice_number, child_req.total_price))
                else:
                    if hasattr(req, 'payment'): _add_file_to_zip(
                        zip_file, req.payment, 'payment_photo',
                        _build_attachment_name(req.total_price, '支付凭证')
                    )
                    if hasattr(req, 'purchase_order'): _add_file_to_zip(
                        zip_file, req.purchase_order, 'po_form',
                        _build_attachment_name(req.total_price, '请购单')
                    )
                    if hasattr(req, 'contract'): _add_file_to_zip(
                        zip_file, req.contract, 'contract_form',
                        _build_attachment_name(req.total_price, '合同')
                    )
                    if hasattr(req, 'acceptance'): _add_file_to_zip(
                        zip_file, req.acceptance, 'acceptance_photo',
                        _build_attachment_name(req.total_price, '收货凭证')
                    )
                    _add_invoices_to_zip(zip_file, req.invoices.all(), req.total_price)
                    for invoice in req.invoices.all():
                        invoice_rows.append((invoice.invoice_number, req.total_price))
                    if hasattr(req, 'inspection'): _add_file_to_zip(
                        zip_file, req.inspection, 'inspection_photo',
                        _build_attachment_name(req.total_price, '验收凭证')
                    )

            try:
                excel_name, excel_bytes = _build_invoice_excel(invoice_rows, selected_total_sum)
                zip_file.writestr(excel_name, excel_bytes)
            except Exception as e:
                return Response({"error": f"生成发票信息表失败: {str(e)}"}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

        buffer.seek(0)
        response = HttpResponse(buffer, content_type='application/zip')
        response['Content-Disposition'] = 'attachment; filename="采购报销附件.zip"'
        return response


class MergeRequestsView(APIView):
    permission_classes = [IsAuthenticated, IsApprover]

    @transaction.atomic
    def post(self, request, *args, **kwargs):
        request_ids = request.data.get('request_ids', [])
        if not isinstance(request_ids, list) or len(request_ids) < 2:
            return Response({"error": "请至少选择 2 个采购申请进行合并。"}, status=status.HTTP_400_BAD_REQUEST)

        children_requests_qs = PurchaseRequest.objects.filter(id__in=request_ids)
        for req in children_requests_qs:
            self.check_object_permissions(request, req)

        children_requests = list(children_requests_qs)
        if len(children_requests) != len(request_ids):
            return Response({"error": "部分选择的申请不存在或您无权操作。"}, status=status.HTTP_400_BAD_REQUEST)

        first_req = children_requests[0]
        expense_type = first_req.expense_type
        statuses = set()
        for req in children_requests:
            statuses.add(req.status)
            if req.parent_request is not None:
                return Response({"error": f"申请 '{req.order_number}' 已经被合并过，无法再次合并。"},
                                status=status.HTTP_400_BAD_REQUEST)

            if req.expense_type != expense_type:
                return Response({"error": "只能合并来自同一经费类型的申请。"},
                                status=status.HTTP_400_BAD_REQUEST)

        inspection_stage = {'invoiced'}
        reimbursement_stage = {'accepted', 'inspection_skipped'}
        if statuses.issubset(inspection_stage):
            parent_status = 'invoiced'
        elif statuses.issubset(reimbursement_stage):
            parent_status = 'accepted' if 'accepted' in statuses else 'inspection_skipped'
        else:
            return Response({"error": "只能合并同一阶段的申请（待验收或待报销）。"},
                            status=status.HTTP_400_BAD_REQUEST)

        handler_ids = [req.handler_id for req in children_requests if req.handler_id]
        unique_handler_ids = list(dict.fromkeys(handler_ids))
        if unique_handler_ids:
            resolved_handler_id = request.user.id if request.user.id in unique_handler_ids else unique_handler_ids[0]
        else:
            resolved_handler_id = request.user.id

        approver_ids = [req.approved_by_id for req in children_requests if req.approved_by_id]
        resolved_approved_by_id = approver_ids[0] if approver_ids else None

        total_price = sum(req.total_price for req in children_requests)

        parent_request = PurchaseRequest.objects.create(
            applicant=request.user, status=parent_status, expense_type=expense_type,
            platform="多个平台", total_price=total_price, applicant_tutor=first_req.applicant.assigned_tutor,
            handler_id=resolved_handler_id, approved_by_id=resolved_approved_by_id
        )

        for req in children_requests:
            RequestItem.objects.filter(purchase_request=req).update(
                purchase_request=parent_request,
                source_request=req
            )
            req.merged_from_status = req.status
            req.status = 'merged'
            req.parent_request = parent_request
            req.save(update_fields=['merged_from_status', 'status', 'parent_request'])

        parent_request.order_number = "MERGED-" + "-".join(
            [req.order_number for req in children_requests if req.order_number])[:40]
        parent_request.save(update_fields=['order_number'])

        serializer = PurchaseRequestSerializer(parent_request)
        return Response(serializer.data, status=status.HTTP_201_CREATED)


class UnmergeRequestView(APIView):
    permission_classes = [IsAuthenticated, IsApprover]

    @transaction.atomic
    def post(self, request, pk, *args, **kwargs):
        try:
            parent_request = PurchaseRequest.objects.select_for_update().prefetch_related(
                'merged_children', 'items'
            ).get(pk=pk)
        except PurchaseRequest.DoesNotExist:
            raise Http404

        if not parent_request.merged_children.exists():
            return Response({"error": "该申请不是合并单或没有可撤销的子申请。"}, status=status.HTTP_400_BAD_REQUEST)

        if parent_request.status not in {'accepted', 'inspection_skipped'}:
            return Response({"error": "仅允许在“待报销”阶段撤销合并。"}, status=status.HTTP_400_BAD_REQUEST)

        children = list(parent_request.merged_children.all())
        child_ids = {child.id for child in children}

        for child in children:
            if not child.merged_from_status:
                return Response(
                    {"error": f"子申请 {child.order_number or child.id} 缺少合并前状态记录，无法精准撤销。"},
                    status=status.HTTP_400_BAD_REQUEST
                )

        items = list(RequestItem.objects.filter(purchase_request=parent_request))
        for item in items:
            if not item.source_request_id:
                return Response(
                    {"error": "合并单存在未记录来源的物品，无法精准撤销。"},
                    status=status.HTTP_400_BAD_REQUEST
                )
            if item.source_request_id not in child_ids:
                return Response(
                    {"error": "合并单物品来源不属于当前子申请，无法精准撤销。"},
                    status=status.HTTP_400_BAD_REQUEST
                )

        for child in children:
            RequestItem.objects.filter(
                purchase_request=parent_request,
                source_request=child
            ).update(
                purchase_request=child,
                source_request=None
            )

        if RequestItem.objects.filter(purchase_request=parent_request).exists():
            transaction.set_rollback(True)
            return Response({"error": "仍有物品未能还原，已中止撤销。"}, status=status.HTTP_400_BAD_REQUEST)

        for child in children:
            child.status = child.merged_from_status
            child.merged_from_status = None
            child.parent_request = None
            child.save(update_fields=['status', 'merged_from_status', 'parent_request'])

        parent_request.delete()
        return Response({"status": "合并已撤销。"}, status=status.HTTP_200_OK)


class PackageAttachmentsView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, pk, *args, **kwargs):
        try:
            req = PurchaseRequest.objects.prefetch_related(
                'merged_children__acceptance', 'merged_children__purchase_order',
                'merged_children__contract', 'merged_children__invoices', 'merged_children__inspection',
                'invoices'
            ).select_related(
                'acceptance', 'purchase_order', 'contract', 'inspection'
            ).get(pk=pk)
        except PurchaseRequest.DoesNotExist:
            raise Http404

        is_approver = IsApprover().has_object_permission(request, self, req)
        is_owner = request.user == req.applicant
        is_child_owner = False
        if req.merged_children.exists():
            if request.user.id in req.merged_children.values_list('applicant_id', flat=True):
                is_child_owner = True

        if not (is_owner or is_child_owner or request.user.is_staff or is_approver):
            return Response({"error": "您没有权限下载此申请的附件。"}, status=status.HTTP_403_FORBIDDEN)

        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            if req.merged_children.exists():
                for child in req.merged_children.all():
                    child_prefix = child.order_number or f'child_{child.id}'
                    _add_file_to_zip(zip_file, getattr(child, 'acceptance', None), 'acceptance_photo',
                                     f"{child_prefix}-收货照片")
                    _add_file_to_zip(zip_file, getattr(child, 'purchase_order', None), 'po_form',
                                     f"{child_prefix}-请购单")
                    _add_file_to_zip(zip_file, getattr(child, 'contract', None), 'contract_form',
                                     f"{child_prefix}-合同文件")
                    _add_invoices_to_zip(zip_file, child.invoices.all(), f"{child_prefix}-发票截图")
                    _add_file_to_zip(zip_file, getattr(child, 'inspection', None), 'inspection_photo',
                                     f"{child_prefix}-验收照片")
            else:
                prefix = req.order_number or f'request_{req.id}'
                _add_file_to_zip(zip_file, getattr(req, 'acceptance', None), 'acceptance_photo', f"{prefix}-收货照片")
                _add_file_to_zip(zip_file, getattr(req, 'purchase_order', None), 'po_form', f"{prefix}-请购单")
                _add_file_to_zip(zip_file, getattr(req, 'contract', None), 'contract_form', f"{prefix}-合同文件")
                _add_invoices_to_zip(zip_file, req.invoices.all(), f"{prefix}-发票截图")
                _add_file_to_zip(zip_file, getattr(req, 'inspection', None), 'inspection_photo', f"{prefix}-验收照片")

        buffer.seek(0)
        zip_filename = f"{req.order_number or f'request_{req.id}'}_附件.zip"
        zip_filename_encoded = zip_filename.encode('utf-8').decode('latin-1')
        response = HttpResponse(buffer, content_type='application/zip')
        response['Content-Disposition'] = f'attachment; filename="{zip_filename_encoded}"'
        return response


class PackageReimbursementImagesView(APIView):
    """
    打包指定采购申请的报销相关图片（仅图片类附件）。
    包含：支付截图、收货照片、发票截图、验收照片、报销单照片。
    """
    permission_classes = [IsAuthenticated]

    def get(self, request, pk, *args, **kwargs):
        try:
            req = PurchaseRequest.objects.prefetch_related(
                'merged_children__payment', 'merged_children__acceptance', 'merged_children__invoices',
                'merged_children__inspection', 'merged_children__reimbursement_details',
                'invoices', 'reimbursement_details'
            ).select_related(
                'payment', 'acceptance', 'inspection', 'reimbursement_details'
            ).get(pk=pk)
        except PurchaseRequest.DoesNotExist:
            raise Http404

        is_approver = IsApprover().has_object_permission(request, self, req)
        is_owner = request.user == req.applicant
        is_child_owner = False
        if req.merged_children.exists():
            if request.user.id in req.merged_children.values_list('applicant_id', flat=True):
                is_child_owner = True

        if not (is_owner or is_child_owner or request.user.is_staff or is_approver):
            return Response({"error": "您没有权限下载此申请的附件。"}, status=status.HTTP_403_FORBIDDEN)

        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            if req.merged_children.exists():
                # 父申请可能有共用验收/报销图片
                if hasattr(req, 'inspection'):
                    _add_file_to_zip(
                        zip_file, req.inspection, 'inspection_photo',
                        _build_attachment_name(req.total_price, '验收凭证')
                    )
                if hasattr(req, 'reimbursement_details'):
                    _add_reimbursement_to_zip(zip_file, req.reimbursement_details, req.total_price)

                for child_req in req.merged_children.all():
                    if hasattr(child_req, 'payment'):
                        _add_file_to_zip(
                            zip_file, child_req.payment, 'payment_photo',
                            _build_attachment_name(child_req.total_price, '支付凭证')
                        )
                    if hasattr(child_req, 'acceptance'):
                        _add_file_to_zip(
                            zip_file, child_req.acceptance, 'acceptance_photo',
                            _build_attachment_name(child_req.total_price, '收货凭证')
                        )
                    _add_invoices_to_zip(zip_file, child_req.invoices.all(), child_req.total_price)
                    if hasattr(child_req, 'inspection'):
                        _add_file_to_zip(
                            zip_file, child_req.inspection, 'inspection_photo',
                            _build_attachment_name(child_req.total_price, '验收凭证')
                        )
                    if hasattr(child_req, 'reimbursement_details'):
                        _add_reimbursement_to_zip(zip_file, child_req.reimbursement_details, child_req.total_price)
            else:
                if hasattr(req, 'payment'):
                    _add_file_to_zip(
                        zip_file, req.payment, 'payment_photo',
                        _build_attachment_name(req.total_price, '支付凭证')
                    )
                if hasattr(req, 'acceptance'):
                    _add_file_to_zip(
                        zip_file, req.acceptance, 'acceptance_photo',
                        _build_attachment_name(req.total_price, '收货凭证')
                    )
                _add_invoices_to_zip(zip_file, req.invoices.all(), req.total_price)
                if hasattr(req, 'inspection'):
                    _add_file_to_zip(
                        zip_file, req.inspection, 'inspection_photo',
                        _build_attachment_name(req.total_price, '验收凭证')
                    )
                if hasattr(req, 'reimbursement_details'):
                    _add_reimbursement_to_zip(zip_file, req.reimbursement_details, req.total_price)

        buffer.seek(0)
        zip_filename = f"{req.order_number or f'request_{req.id}'}_报销图片.zip"
        zip_filename_encoded = zip_filename.encode('utf-8').decode('latin-1')
        response = HttpResponse(buffer, content_type='application/zip')
        response['Content-Disposition'] = f'attachment; filename=\"{zip_filename_encoded}\"'
        return response


class ExportC2CPlatformRecordsView(APIView):
    """
    导出公对公采购流程总览记录。
    - request_ids: 导出当前勾选记录
    - platform: 未提供 request_ids 时，导出该平台全部可见记录
    """
    permission_classes = [IsAuthenticated]

    def _visible_queryset(self, request):
        user = request.user
        queryset = PurchaseRequest.objects.filter(
            expense_type='c2c',
            parent_request__isnull=True,
        )

        user_roles = set(user.roles.values_list('name', flat=True))
        if '系统管理员' in user_roles or '导师用户' in user_roles:
            visible_queryset = queryset
        else:
            managed_platforms = list(
                Platform.objects.filter(managers=user, category='c2c').values_list('name', flat=True)
            )
            if managed_platforms:
                visible_queryset = queryset.filter(platform__in=managed_platforms)
            else:
                visible_queryset = queryset.none()

        return visible_queryset.select_related(
            'applicant', 'applicant__assigned_tutor', 'approved_by', 'handler',
            'payment', 'acceptance', 'inspection', 'purchase_order', 'contract', 'reimbursement_details',
        ).prefetch_related(
            'items', 'invoices',
            'merged_children__items',
            'merged_children__applicant', 'merged_children__applicant__assigned_tutor',
            'merged_children__approved_by', 'merged_children__handler',
            'merged_children__payment', 'merged_children__acceptance', 'merged_children__inspection',
            'merged_children__purchase_order', 'merged_children__contract',
            'merged_children__invoices', 'merged_children__reimbursement_details',
        ).distinct()

    def post(self, request, *args, **kwargs):
        raw_ids = request.data.get('request_ids') or []
        platform_name = (request.data.get('platform') or '').strip()
        queryset = self._visible_queryset(request)

        if raw_ids:
            if not isinstance(raw_ids, list):
                return Response({"error": "request_ids 必须是数组。"}, status=status.HTTP_400_BAD_REQUEST)

            try:
                request_ids = sorted({int(item) for item in raw_ids})
            except (TypeError, ValueError):
                return Response({"error": "request_ids 包含无效ID。"}, status=status.HTTP_400_BAD_REQUEST)

            queryset = queryset.filter(id__in=request_ids)
            if queryset.count() != len(request_ids):
                return Response({"error": "部分勾选记录不存在或无权导出。"}, status=status.HTTP_403_FORBIDDEN)

            platforms = list(queryset.values_list('platform', flat=True).distinct())
            export_platform_name = platforms[0] if len(platforms) == 1 else '选中'
        else:
            if not platform_name:
                return Response({"error": "请选择平台，或先勾选要导出的记录。"}, status=status.HTTP_400_BAD_REQUEST)
            queryset = queryset.filter(platform=platform_name)
            export_platform_name = platform_name

        requests_to_export = list(queryset.order_by('request_date', 'order_number', 'id'))
        if not requests_to_export:
            return Response({"error": "没有可导出的记录。"}, status=status.HTTP_400_BAD_REQUEST)

        filename, excel_bytes = _build_c2c_records_excel(requests_to_export, export_platform_name)
        response = HttpResponse(
            excel_bytes,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        ascii_name = f"c2c_purchase_records_{timezone.localtime().strftime('%Y%m%d')}.xlsx"
        response['Content-Disposition'] = f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"
        return response


# --- 【新功能】 ---
class BulkDeleteRequestsView(APIView):
    """
    供系统管理员批量删除采购申请的视图。
    """
    permission_classes = [IsAuthenticated, IsSystemAdmin]

    @transaction.atomic
    def post(self, request, *args, **kwargs):
        request_ids = request.data.get('request_ids', [])
        if not isinstance(request_ids, list) or not request_ids:
            return Response({"error": "请提供需要删除的申请ID列表。"}, status=status.HTTP_400_BAD_REQUEST)

        # 删除父申请会通过 CASCADE 自动删除所有关联的对象（items, invoice, etc.）
        # 同时，也需要删除可能被这些父申请合并的子申请
        parent_requests = PurchaseRequest.objects.filter(id__in=request_ids)

        # 找到所有即将被删除的父申请所关联的子申请
        children_to_delete = PurchaseRequest.objects.filter(parent_request__in=parent_requests)

        # 合并ID并去重，以防万一
        all_ids_to_delete = set(request_ids) | set(children_to_delete.values_list('id', flat=True))

        # 执行删除
        deleted_count, _ = PurchaseRequest.objects.filter(id__in=all_ids_to_delete).delete()

        return Response({"status": f"成功删除了 {deleted_count} 个相关申请记录。"}, status=status.HTTP_200_OK)
# --- 【新功能结束】 ---
